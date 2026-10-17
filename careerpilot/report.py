"""Export the preparation guide (a plain dict from agent.py) to a Word document."""
import io

from docx import Document


def build_guide_docx(guide):
    """Return the guide as DOCX bytes. Raises ValueError if guide is empty."""
    if not guide:
        raise ValueError("There is no guide to export.")

    doc = Document()
    doc.add_heading("Interview Preparation Guide", 0)
    doc.add_paragraph(f"Generated on {guide['generated_on']} by CareerPilot Mini")

    doc.add_heading("Target job summary", 1)
    doc.add_paragraph(f"Role: {guide['job_title']}")
    doc.add_paragraph(guide["job_summary"])

    doc.add_heading("Skills matched against your resume", 1)
    if guide["matched_skills"]:
        for s in guide["matched_skills"]:
            doc.add_paragraph(s, style="List Bullet")
    else:
        doc.add_paragraph("No job skills were found in the resume text.")

    doc.add_heading("Skills not evidenced in your resume", 1)
    if guide["missing_skills"]:
        for s in guide["missing_skills"]:
            doc.add_paragraph(s, style="List Bullet")
    else:
        doc.add_paragraph("Every skill detected in the job description appears in the resume.")

    res_by_id = {r["id"]: r for r in guide["resources"]}
    doc.add_heading("Prioritized interview preparation topics", 1)
    for i, t in enumerate(guide["topics"], 1):
        doc.add_paragraph().add_run(f"{i}. {t['topic']}").bold = True
        if t["reason"]:
            doc.add_paragraph(t["reason"])
        for rid in t["resource_ids"]:
            r = res_by_id.get(rid)
            if r:
                doc.add_paragraph(f"See resource [{rid}]: {r['title']}", style="List Bullet")

    doc.add_heading("Preparation strategy", 1)
    for step in guide["strategy"]:
        doc.add_paragraph(step, style="List Bullet")

    doc.add_heading("Practice questions (generated, not sourced)", 1)
    if guide["practice_questions"]:
        for q in guide["practice_questions"]:
            doc.add_paragraph(f"[{q['topic']}] {q['question']}", style="List Bullet")
    else:
        doc.add_paragraph("No practice questions were generated (the LLM step did not run).")

    doc.add_heading("Research resources", 1)
    if guide["resources"]:
        for r in guide["resources"]:
            doc.add_paragraph(f"[{r['id']}] {r['title']}")
            doc.add_paragraph(f"Topic: {r['topic']}")
            doc.add_paragraph(f"URL: {r['url']}")
            published = r["published"] or "not verified"
            doc.add_paragraph(f"Published: {published} ({r['date_note']})")
    else:
        doc.add_paragraph("No resources were retrieved. Live research was unavailable or turned off.")

    doc.add_heading("Limitations and uncertainty notes", 1)
    for note in guide["warnings"] + guide["limitations"]:
        doc.add_paragraph(note, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
